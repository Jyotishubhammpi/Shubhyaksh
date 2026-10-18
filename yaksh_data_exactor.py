import re
import os
import time
import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Alignment, Font
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

# Configure Chrome options to maintain your login session
chrome_options = Options()
chrome_options = Options()
chrome_options.add_argument("--headless")              # बिना स्क्रीन बैकग्राउंड में चलाने के लिए
chrome_options.add_argument("--no-sandbox")             # Codespaces/Linux के लिए ज़रूरी
chrome_options.add_argument("--disable-dev-shm-usage")  # मेमोरी क्रैश से बचने के लिए
driver = webdriver.Chrome(options=chrome_options)
custom_profile_path = r"C:\Users\HP\Desktop\MySeleniumProfile"
chrome_options.add_argument(f"--user-data-dir={custom_profile_path}")
chrome_options.add_argument("--disable-blink-features=AutomationControlled")
chrome_options.add_experimental_option("excludeSwitches", ["enable-automation"])
chrome_options.add_experimental_option('useAutomationExtension', False)

# 🟢 सभी सर्कल्स और उनके थाने
circle_police_stations = {
    "CIRCLE AKBARPUR": ["AKBARPUR", "RURA", "GAJNER", "RANIYA"],
    "CIRCLE BHOGNIPUR": ["BHOGNIPUR", "MOOSANAGAR", "DEVRAHAT", "SATTI"],
    "CIRCLE SIKANDARA": ["SIKANDRA", "RAJPUR", "AMRAHAT"],
    "CIRCLE DERAPUR": ["DERAPUR", "MANGALPUR", "BARAUR"],
    "CIRCLE RASOOLABAD": ["RASOOLABAD", "SHIVLI"],
}

driver = None

try:
    driver = webdriver.Chrome(options=chrome_options)
    wait = WebDriverWait(driver, 25)

    driver.get("https://yaksh.ai/LOGIN")
    print("Loading page...")

    # Step 1: Click on "Criminal Record of My Jurisdiction"
    step1_elem = wait.until(EC.element_to_be_clickable((By.XPATH, "//p[contains(text(), 'Criminal Record of My Jurisdiction')]")))
    driver.execute_script("arguments[0].click();", step1_elem)
    print("Step 1 completed.")
    time.sleep(2)

    # Step 2: Click on the 'OPEN' button under "Based on Criminal Address"
    try:
        step2_elem = wait.until(EC.presence_of_element_located((By.XPATH, "//p[contains(text(), 'Based on Criminal Address')]/ancestor::div[contains(@class, 'd-flex')]//p[text()='OPEN']")))
    except:
        step2_elem = wait.until(EC.presence_of_element_located((By.XPATH, "(//p[text()='OPEN'])[1]")))
    
    driver.execute_script("arguments[0].click();", step2_elem)
    print("Step 2 completed.")
    time.sleep(3)

    # -------------------------------------------------------------------------
    # PART 1: Fetching Total Criminal Records
    # -------------------------------------------------------------------------
    print("\n--- [Phase 1/3] Fetching Total Criminal Records ---")
    total_records_dict = {}

    for circle, police_stations in circle_police_stations.items():
        print(f"\nSelecting Circle: {circle}")
        
        circle_input = wait.until(EC.element_to_be_clickable((By.ID, "rc_select_1")))
        driver.execute_script("arguments[0].click();", circle_input)
        time.sleep(1)
        circle_input.send_keys(Keys.CONTROL + "a")
        circle_input.send_keys(Keys.BACKSPACE)
        circle_input.send_keys(circle)
        time.sleep(1.5)

        try:
            circle_option = wait.until(EC.element_to_be_clickable((By.XPATH, f"//div[contains(@class, 'ant-select-item-option') and contains(., '{circle}')]")))
            driver.execute_script("arguments[0].click();", circle_option)
        except:
            circle_input.send_keys(Keys.ENTER)
        
        time.sleep(3.5)

        for ps in police_stations:
            print(f"Processing (Total): Circle - {circle}, Police Station - {ps}")

            ps_input = wait.until(EC.element_to_be_clickable((By.ID, "rc_select_2")))
            driver.execute_script("arguments[0].click();", ps_input)
            time.sleep(1)
            
            ps_input.send_keys(Keys.CONTROL + "a")
            ps_input.send_keys(Keys.BACKSPACE)
            time.sleep(1)
            
            ps_input.send_keys(ps)
            time.sleep(2)

            try:
                ps_option = wait.until(EC.element_to_be_clickable((By.XPATH, f"//div[contains(@class, 'ant-select-item-option') and contains(., '{ps}')]")))
                driver.execute_script("arguments[0].click();", ps_option)
            except:
                ps_input.send_keys(Keys.ENTER)

            time.sleep(2)

            old_text = ""
            try:
                old_elem = driver.find_element(By.XPATH, "//p[contains(text(), 'Criminal Record Found') or contains(@class, 'fs-18')]")
                old_text = old_elem.text.strip()
            except:
                pass

            apply_btn = wait.until(EC.element_to_be_clickable((By.XPATH, "//button[@type='submit' and contains(., 'Apply')]")))
            driver.execute_script("arguments[0].click();", apply_btn)

            count = 0
            start_time = time.time()
            while True:
                try:
                    record_elem = driver.find_element(By.XPATH, "//p[contains(text(), 'Criminal Record Found') or contains(@class, 'fs-18')]")
                    txt = record_elem.text.strip()
                    
                    if txt and txt != old_text and any(char.isdigit() for char in txt):
                        clean_txt = txt.replace(",", "")
                        count_match = re.search(r'\d+', clean_txt)
                        if count_match:
                            found_val = int(count_match.group(0))
                            count = found_val
                            break
                except:
                    pass
                
                if time.time() - start_time > 15:
                    try:
                        record_elem = driver.find_element(By.XPATH, "//p[contains(text(), 'Criminal Record Found') or contains(@class, 'fs-18')]")
                        clean_txt = record_elem.text.replace(",", "")
                        count_match = re.search(r'\d+', clean_txt)
                        if count_match:
                            count = int(count_match.group(0))
                    except:
                        pass
                    break
                time.sleep(0.5)

            total_records_dict[(circle, ps)] = count
            print(f"-> {circle} | {ps} Total Criminals: {count}")
            time.sleep(1.5)

    # -------------------------------------------------------------------------
    # PART 2: Fetching Beat Verification Completed Records
    # -------------------------------------------------------------------------
    print("\n--- [Phase 2/3] Fetching Beat Verification Completed Records ---")
    
    try:
        verified_tab = wait.until(EC.element_to_be_clickable((By.XPATH, "//div[contains(@id, 'tab-verified') or contains(text(), 'Beat Verification Completed')]")))
        driver.execute_script("arguments[0].click();", verified_tab)
        print("Switched to 'Beat Verification Completed' tab successfully!")
        time.sleep(3)
    except Exception as e:
        print(f"Note regarding tab selection: {e}")

    verified_records_dict = {}

    for circle, police_stations in circle_police_stations.items():
        try:
            circle_input = wait.until(EC.element_to_be_clickable((By.ID, "rc_select_1")))
            driver.execute_script("arguments[0].click();", circle_input)
            time.sleep(1)
            circle_input.send_keys(Keys.CONTROL + "a")
            circle_input.send_keys(Keys.BACKSPACE)
            circle_input.send_keys(circle)
            time.sleep(1.5)

            circle_option = wait.until(EC.element_to_be_clickable((By.XPATH, f"//div[contains(@class, 'ant-select-item-option') and contains(., '{circle}')]")))
            driver.execute_script("arguments[0].click();", circle_option)
            time.sleep(3.5)
        except:
            pass

        for ps in police_stations:
            print(f"Processing (Verified): Circle - {circle}, Police Station - {ps}")

            ps_input = wait.until(EC.element_to_be_clickable((By.ID, "rc_select_2")))
            driver.execute_script("arguments[0].click();", ps_input)
            time.sleep(1)
            
            ps_input.send_keys(Keys.CONTROL + "a")
            ps_input.send_keys(Keys.BACKSPACE)
            time.sleep(1)
            
            ps_input.send_keys(ps)
            time.sleep(2)

            try:
                ps_option = wait.until(EC.element_to_be_clickable((By.XPATH, f"//div[contains(@class, 'ant-select-item-option') and contains(., '{ps}')]")))
                driver.execute_script("arguments[0].click();", ps_option)
            except:
                ps_input.send_keys(Keys.ENTER)

            time.sleep(2)

            old_text = ""
            try:
                old_elem = driver.find_element(By.XPATH, "//p[contains(text(), 'Record Found') or contains(@class, 'fs-18')]")
                old_text = old_elem.text.strip()
            except:
                pass

            apply_btn = wait.until(EC.element_to_be_clickable((By.XPATH, "//button[@type='submit' and contains(., 'Apply')]")))
            driver.execute_script("arguments[0].click();", apply_btn)

            count = 0
            start_time = time.time()
            while True:
                try:
                    record_elem = driver.find_element(By.XPATH, "//p[contains(text(), 'Record Found') or contains(@class, 'fs-18')]")
                    txt = record_elem.text.strip()
                    
                    if txt and txt != old_text and any(char.isdigit() for char in txt):
                        clean_txt = txt.replace(",", "")
                        count_match = re.search(r'\d+', clean_txt)
                        if count_match:
                            found_val = int(count_match.group(0))
                            count = found_val
                            break
                except:
                    pass
                
                if time.time() - start_time > 15:
                    try:
                        record_elem = driver.find_element(By.XPATH, "//p[contains(text(), 'Record Found') or contains(@class, 'fs-18')]")
                        clean_txt = record_elem.text.replace(",", "")
                        count_match = re.search(r'\d+', clean_txt)
                        if count_match:
                            count = int(count_match.group(0))
                    except:
                        pass
                    break
                time.sleep(0.5)

            verified_records_dict[(circle, ps)] = count
            print(f"-> {circle} | {ps} Beat Verification Completed: {count}")
            time.sleep(1.5)

    # -------------------------------------------------------------------------
    # PART 3: Fetching Incorrect Address Reported by Beat Officer Records
    # -------------------------------------------------------------------------
    print("\n--- [Phase 3/3] Fetching Incorrect Address Records ---")
    
    try:
        incorrect_tab = wait.until(EC.element_to_be_clickable((By.XPATH, "//div[contains(@id, 'tab-incorrect_address') or contains(text(), 'Incorrect Address Reported by Beat Officer')]")))
        driver.execute_script("arguments[0].click();", incorrect_tab)
        print("Switched to 'Incorrect Address' tab successfully!")
        time.sleep(3)
    except Exception as e:
        print(f"Note regarding tab selection: {e}")

    incorrect_records_dict = {}

    for circle, police_stations in circle_police_stations.items():
        try:
            circle_input = wait.until(EC.element_to_be_clickable((By.ID, "rc_select_1")))
            driver.execute_script("arguments[0].click();", circle_input)
            time.sleep(1)
            circle_input.send_keys(Keys.CONTROL + "a")
            circle_input.send_keys(Keys.BACKSPACE)
            circle_input.send_keys(circle)
            time.sleep(1.5)

            circle_option = wait.until(EC.element_to_be_clickable((By.XPATH, f"//div[contains(@class, 'ant-select-item-option') and contains(., '{circle}')]")))
            driver.execute_script("arguments[0].click();", circle_option)
            time.sleep(3.5)
        except:
            pass

        for ps in police_stations:
            print(f"Processing (Incorrect Address): Circle - {circle}, Police Station - {ps}")

            ps_input = wait.until(EC.element_to_be_clickable((By.ID, "rc_select_2")))
            driver.execute_script("arguments[0].click();", ps_input)
            time.sleep(1)
            
            ps_input.send_keys(Keys.CONTROL + "a")
            ps_input.send_keys(Keys.BACKSPACE)
            time.sleep(1)
            
            ps_input.send_keys(ps)
            time.sleep(2)

            try:
                ps_option = wait.until(EC.element_to_be_clickable((By.XPATH, f"//div[contains(@class, 'ant-select-item-option') and contains(., '{ps}')]")))
                driver.execute_script("arguments[0].click();", ps_option)
            except:
                ps_input.send_keys(Keys.ENTER)

            time.sleep(2)

            old_text = ""
            try:
                old_elem = driver.find_element(By.XPATH, "//p[contains(text(), 'Record Found') or contains(@class, 'fs-18')]")
                old_text = old_elem.text.strip()
            except:
                pass

            apply_btn = wait.until(EC.element_to_be_clickable((By.XPATH, "//button[@type='submit' and contains(., 'Apply')]")))
            driver.execute_script("arguments[0].click();", apply_btn)

            count = 0
            start_time = time.time()
            while True:
                try:
                    record_elem = driver.find_element(By.XPATH, "//p[contains(text(), 'Record Found') or contains(@class, 'fs-18')]")
                    txt = record_elem.text.strip()
                    
                    if txt and txt != old_text and any(char.isdigit() for char in txt):
                        clean_txt = txt.replace(",", "")
                        count_match = re.search(r'\d+', clean_txt)
                        if count_match:
                            found_val = int(count_match.group(0))
                            count = found_val
                            break
                except:
                    count = 0
                    break
                
                if time.time() - start_time > 15:
                    try:
                        record_elem = driver.find_element(By.XPATH, "//p[contains(text(), 'Record Found') or contains(@class, 'fs-18')]")
                        clean_txt = record_elem.text.replace(",", "")
                        count_match = re.search(r'\d+', clean_txt)
                        if count_match:
                            count = int(count_match.group(0))
                    except:
                        count = 0
                    break
                time.sleep(0.5)

            incorrect_records_dict[(circle, ps)] = count
            print(f"-> {circle} | {ps} Incorrect Address: {count}")
            time.sleep(1.5)

    # -------------------------------------------------------------------------
    # PART 4: Save Data into Excel File with Numeric Values, Center Alignment & Total Row
    # -------------------------------------------------------------------------
    final_data_list = []
    serial_no = 1

    for circle, police_stations in circle_police_stations.items():
        for ps in police_stations:
            final_data_list.append({
                "S.No.": serial_no,
                "Circle Name": circle,
                "Police Station": ps,
                "Total Criminal": total_records_dict.get((circle, ps), 0),
                "Beat Verification Completed": verified_records_dict.get((circle, ps), 0),
                "Incorrect Address": incorrect_records_dict.get((circle, ps), 0)
            })
            serial_no += 1

    # Desktop पर 'YAKSH' फ़ोल्डर बनाना
    desktop_path = os.path.join(os.path.expanduser("~"), "Desktop")
    yaksh_folder = os.path.join(desktop_path, "YAKSH")
    
    if not os.path.exists(yaksh_folder):
        os.makedirs(yaksh_folder)

    # Date-wise Excel File का पाथ बनाना
    current_date = datetime.datetime.now().strftime("%Y-%m-%d")
    excel_filename = os.path.join(yaksh_folder, f"{current_date}_criminal_records.xlsx")
    sheet_name = "Beat Verification"

    # Pandas DataFrame का उपयोग करना
    df = pd.DataFrame(final_data_list)
    with pd.ExcelWriter(excel_filename, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)

    # openpyxl के ज़रिए Numeric फॉर्मेट, Center Alignment और Total Row जोड़ना
    wb = openpyxl.load_workbook(excel_filename)
    ws = wb[sheet_name]

    center_alignment = Alignment(horizontal="center", vertical="center")
    bold_font = Font(bold=True)

    max_row = ws.max_row
    
    # 1. सभी डेटा सेल्स को सेंटर अलाइन करना और न्यूमेरिक (integer) बनाना
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.alignment = center_alignment
            # कॉलम 4, 5, 6 (Total Criminal, Beat Verification, Incorrect Address) को प्रॉपर नंबर बनाना
            if cell.row > 1 and cell.column in [4, 5, 6]:
                try:
                    cell.value = int(cell.value)
                except:
                    pass

    # 2. सबसे नीचे 'Total' की Row जोड़ना और SUM फॉर्मूला लगाना
    total_row_idx = max_row + 1
    ws.cell(row=total_row_idx, column=3, value="Total").font = bold_font
    ws.cell(row=total_row_idx, column=3).alignment = center_alignment

    # Total Criminal (Column D) का फॉर्मूला
    cell_d = ws.cell(row=total_row_idx, column=4, value=f"=SUM(D2:D{max_row})")
    cell_d.font = bold_font
    cell_d.alignment = center_alignment

    # Beat Verification Completed (Column E) का फॉर्मूला
    cell_e = ws.cell(row=total_row_idx, column=5, value=f"=SUM(E2:E{max_row})")
    cell_e.font = bold_font
    cell_e.alignment = center_alignment

    # Incorrect Address (Column F) का फॉर्मूला
    cell_f = ws.cell(row=total_row_idx, column=6, value=f"=SUM(F2:F{max_row})")
    cell_f.font = bold_font
    cell_f.alignment = center_alignment

    wb.save(excel_filename)

    print(f"\nAll data successfully saved to Excel file:\nPath: {excel_filename}\nSheet Name: '{sheet_name}' (Numeric format & Total row added!)")
    input("\nProcess complete! Press Enter to close the browser...")

except Exception as e:
    print(f"\nAn error occurred: {e}")
    input("\nPress Enter to close the browser after checking the error...")

finally:
    if driver:
        driver.quit()
